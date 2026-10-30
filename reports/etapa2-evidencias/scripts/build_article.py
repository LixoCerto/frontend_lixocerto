from pathlib import Path
from docx import Document
from docx.shared import Cm, Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import re

PROJECT = Path('/home/ubuntu/LixoCerto')
FRONT = PROJECT / 'frontend_lixocerto'
EVIDENCE = FRONT / 'reports/etapa2-evidencias'
SOURCE = Path('/home/ubuntu/upload/TCC_LixoCerto_TCCII_escopo_atual.docx')
OUTPUT = PROJECT / 'TCC_LixoCerto_TCCII_Etapa2.docx'

TITLE = ('LIXOCERTO: DESENVOLVIMENTO E AVALIAÇÃO TÉCNICA DE UMA PLATAFORMA WEB '
         'GAMIFICADA PARA ORIENTAÇÃO SOBRE RESÍDUOS SÓLIDOS E ELETROELETRÔNICOS')
ABSTRACT_PT = (
    'No Brasil, foram geradas 81,6 milhões de toneladas de resíduos sólidos urbanos em 2024, e 40,3% tiveram destinação inadequada; em 2022, '
    'o país gerou 2,44 milhões de toneladas de resíduos eletroeletrônicos, das quais apenas 3,2% foram documentadas como formalmente coletadas. '
    'Este trabalho teve como objetivo desenvolver e avaliar tecnicamente a plataforma web gamificada LixoCerto. Orientamos o desenvolvimento pelo '
    'paradigma Design Science Research e implementamos módulos de guia de descarte, pontos de coleta, quiz, desafios e check-in com localização e foto, '
    'usando React, Vite, NestJS, gRPC, PostgreSQL, Redis e armazenamento compatível com S3. Avaliamos o artefato por sete jornadas Playwright, onze testes '
    'de integração e um teste E2E gRPC, inspeção das dez heurísticas de Nielsen, axe-core em sete telas e Lighthouse User Flow em três execuções. As sete '
    'jornadas e os doze testes da API passaram. A inspeção identificou 32 achados heurísticos: um de severidade 1, 28 de severidade 2 e três de severidade 3. '
    'O axe registrou zero violações confirmadas e seis regras incompletas que requerem verificação manual. O LCP mediano de laboratório foi 2,93 s; as medianas '
    'de INP por interação variaram de 58 a 246 ms, com CLS igual a zero. O carregamento inicial de JavaScript em /login caiu 35,4% em gzip após code splitting. '
    'Os resultados demonstram execução dos fluxos especificados e redução relativa da transferência inicial, mas não substituem avaliação com usuários nem dados de campo.'
)
ABSTRACT_EN = (
    'Incorrect disposal of solid and electronic waste continues to pose environmental challenges and limits access to reliable local guidance. '
    'In Brazil, 81.6 million tonnes of municipal solid waste were generated in 2024, and 40.3% received inadequate disposal; in 2022, the country generated '
    '2.44 million tonnes of electronic waste, of which only 3.2% were documented as formally collected. This study aimed to develop and technically evaluate '
    'LixoCerto, a gamified web platform. We followed the Design Science Research paradigm and implemented a disposal guide, collection-point search, quizzes, '
    'challenges, and a check-in flow with location and photographs, using React, Vite, NestJS, gRPC, PostgreSQL, Redis, and S3-compatible storage. We evaluated '
    'the artifact through seven Playwright journeys, eleven integration tests and one gRPC E2E test, an inspection using Nielsen’s ten heuristics, axe-core on seven '
    'screens, and three Lighthouse User Flow runs. All seven journeys and all twelve API integration/E2E tests passed. The heuristic inspection identified 32 findings: '
    'one at severity 1, 28 at severity 2, and three at severity 3. Axe reported no confirmed violations and six incomplete rules requiring manual review. Median laboratory '
    'LCP was 2.93 s; interaction INP medians ranged from 58 to 246 ms, with CLS equal to zero. Initial JavaScript transfer on /login decreased by 35.4% in gzip after code '
    'splitting. The results demonstrate execution of the specified flows and a relative reduction in initial transfer, but they do not replace user evaluation or field data.'
)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def mark_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    tr_pr.append(cant_split)


def set_run_font(run, size=12, bold=False, italic=False, name='Times New Roman', color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_text_paragraph(doc, text, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True,
                       size=12, italic=False, bold=False, before=0, after=6,
                       keep_with_next=False, line_spacing=1.5, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.27) if indent else Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=12, bold=True)
    return p


def add_caption(doc, text, *, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if above else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=9, bold=above)
    return p


def add_source(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=9)
    return p


def add_figure(doc, number, image_path, caption, source, alt, width_inches):
    image_path = Path(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    try:
        shape._inline.docPr.set('descr', alt)
        shape._inline.docPr.set('title', caption)
    except Exception:
        pass
    add_caption(doc, f'Figura {number} – {caption}.')
    add_source(doc, f'Fonte: {source}.')


def add_table(doc, number, title, headers, rows, widths_cm, source, font_size=8.5):
    add_caption(doc, f'Tabela {number} – {title}.', above=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.width = Cm(widths_cm[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, 'D9E2E3')
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                set_run_font(r, size=font_size, bold=True)
    mark_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
            cells[i].width = Cm(widths_cm[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    set_run_font(r, size=font_size)
        prevent_row_split(table.rows[-1])
    add_source(doc, f'Fonte: {source}.')
    return table


def add_reference(doc, lead, title, tail):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(lead)
    set_run_font(r, size=10)
    r = p.add_run(title)
    set_run_font(r, size=10, italic=True)
    if tail:
        r = p.add_run(tail)
        set_run_font(r, size=10)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def set_document_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_author_header(doc):
    p = add_text_paragraph(doc,
        'Arthur Nobrega Martins¹; Hygor Gabriel Alves Montenegro¹; Thiago Rodrigues Medeiros¹',
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, after=2, line_spacing=1.0)
    p = add_text_paragraph(doc,
        '¹ Ciência da Computação – Centro Universitário de João Pessoa (UNIPÊ), João Pessoa – PB, Brasil.',
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, after=2, line_spacing=1.0)
    p = add_text_paragraph(doc,
        'E-mails: nobregaarthur63@gmail.com; ahygor31@gmail.com; southiagorm@gmail.com.',
        align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=10, after=8, line_spacing=1.0)


def main():
    doc = Document(str(SOURCE))
    clear_body(doc)
    set_document_styles(doc)

    # Folha de rosto: telefones foram retirados; o número do orientador não foi fornecido ainda.
    add_text_paragraph(doc, TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, bold=True,
                       before=36, after=32, line_spacing=1.2)
    add_text_paragraph(doc, 'IDENTIFICAÇÃO DOS AUTORES', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, bold=True, before=25, after=20, line_spacing=1.0)
    add_text_paragraph(doc, 'Arthur Nobrega Martins', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, bold=True, after=2, line_spacing=1.0)
    add_text_paragraph(doc, 'E-mail: nobregaarthur63@gmail.com', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, after=14, line_spacing=1.0)
    add_text_paragraph(doc, 'Hygor Gabriel Alves Montenegro', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, bold=True, after=2, line_spacing=1.0)
    add_text_paragraph(doc, 'E-mail: ahygor31@gmail.com', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, after=14, line_spacing=1.0)
    add_text_paragraph(doc, 'Prof. Thiago Rodrigues Medeiros', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, bold=True, after=2, line_spacing=1.0)
    add_text_paragraph(doc, 'E-mail: southiagorm@gmail.com', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, after=18, line_spacing=1.0)
    add_text_paragraph(doc, 'Curso de Ciência da Computação – Centro Universitário de João Pessoa (UNIPÊ)',
                       align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, after=6, line_spacing=1.0)
    add_text_paragraph(doc, 'João Pessoa – PB', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, after=6, line_spacing=1.0)
    add_text_paragraph(doc, 'Artigo de desenvolvimento e avaliação técnica apresentado no âmbito do Trabalho de Conclusão de Curso II.',
                       align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=11, after=10, line_spacing=1.0)
    add_text_paragraph(doc, '2026', align=WD_ALIGN_PARAGRAPH.CENTER,
                       indent=False, size=12, after=0, line_spacing=1.0)
    doc.add_page_break()

    add_text_paragraph(doc, TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12,
                       bold=True, after=10, line_spacing=1.0)
    add_author_header(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run('RESUMO: ')
    set_run_font(r, size=10, bold=True)
    r = p.add_run(ABSTRACT_PT)
    set_run_font(r, size=10)
    add_text_paragraph(doc, 'Palavras-chave: Descarte de resíduos; Educação ambiental; Resíduos eletroeletrônicos; Gamificação; Acessibilidade; Desempenho web.',
                       align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, size=10, after=7, line_spacing=1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run('ABSTRACT: ')
    set_run_font(r, size=10, italic=True, bold=True)
    r = p.add_run(ABSTRACT_EN)
    set_run_font(r, size=10, italic=True)
    add_text_paragraph(doc, 'Keywords: Waste disposal; Environmental education; Electronic waste; Gamification; Accessibility; Web performance.',
                       align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, size=10, italic=True, after=9, line_spacing=1.0)

    add_heading(doc, '1 INTRODUÇÃO', 1)
    add_text_paragraph(doc,
        'A gestão de resíduos sólidos depende de infraestrutura, políticas públicas e informação que permita separar e encaminhar cada material ao fluxo adequado. No Brasil, foram geradas 81,6 milhões de toneladas de resíduos sólidos urbanos em 2024; 40,3% tiveram destinação inadequada (ABREMA, 2025). Esses números evidenciam a dimensão do problema, mas não devem ser confundidos com uma estimativa de descarte incorreto de cada categoria nem com uma medida de comportamento individual.')
    add_text_paragraph(doc,
        'Os resíduos de equipamentos eletroeletrônicos requerem atenção específica por conterem materiais recuperáveis e componentes que demandam manejo próprio. Para 2022, o Brasil registrou geração de aproximadamente 2,44 milhões de toneladas de resíduos eletroeletrônicos e coleta formal documentada de cerca de 79 mil toneladas, aproximadamente 3,2% do total estimado (ITU; UNITAR, 2024). A Política Nacional de Resíduos Sólidos e a regulamentação da logística reversa atribuem responsabilidades aos diversos elos da cadeia; ainda assim, a orientação ao consumidor precisa distinguir os fluxos locais e os pontos que realmente recebem cada material (Brasil, 2010, 2020).')
    add_text_paragraph(doc,
        'A educação ambiental é um processo permanente que articula conhecimentos, valores e participação social (Brasil, 1999). Recursos digitais podem ampliar o acesso a orientações e atividades, embora os resultados dependam de contexto, público e desenho pedagógico (Hajj-Hassan; Chaker; Cederqvist, 2024). Do mesmo modo, pontos, recompensas e desafios são elementos de projeto gamificado, não evidências automáticas de engajamento ou aprendizagem (Deterding et al., 2011; Bai; Hew; Huang, 2020).')
    add_text_paragraph(doc,
        'Neste contexto, desenvolvemos o LixoCerto como plataforma web para reunir orientação de descarte, consulta a pontos de coleta e atividades gamificadas. A pergunta de pesquisa é: como desenvolver e avaliar tecnicamente uma plataforma web gamificada quanto à funcionalidade, à usabilidade por inspeção, à acessibilidade e ao desempenho? O objetivo geral é desenvolver e avaliar tecnicamente a plataforma nesses quatro aspectos. Como contribuições, apresentamos o artefato, a avaliação reproduzível de seus fluxos principais, a comparação controlada do carregamento inicial após code splitting e o registro explícito dos limites de dados e de localização do ponto demonstrado.')

    add_heading(doc, '2 FUNDAMENTAÇÃO TEÓRICA', 1)
    add_heading(doc, '2.1 Resíduos, logística reversa e educação ambiental digital', 2)
    add_text_paragraph(doc,
        'A Política Nacional de Resíduos Sólidos diferencia destinação final ambientalmente adequada — que pode envolver reutilização, reciclagem, recuperação e tratamento — de disposição final dos rejeitos (Brasil, 2010). Logo, orientar a pessoa a localizar um ponto não permite afirmar que o material foi recebido, reciclado ou recuperado. Endereços, horários e categorias aceitas variam por município e operador e podem mudar; por isso, as fichas devem apresentar origem, data de revisão e ressalvas de escopo.')
    add_text_paragraph(doc,
        'A literatura sobre educação ambiental mediada por recursos digitais identifica potencial para apoiar conhecimento e participação, mas também heterogeneidade metodológica e resultados dependentes do contexto (Hajj-Hassan; Chaker; Cederqvist, 2024). Por essa razão, neste estudo tratamos o conteúdo editorial e os módulos educativos como funcionalidades implementadas, sem inferir efeitos de aprendizagem ou mudança de comportamento.')

    add_heading(doc, '2.2 Gamificação e organização das tarefas', 2)
    add_text_paragraph(doc,
        'A gamificação consiste na utilização de elementos associados a jogos em contextos que não são jogos, como pontos, metas, desafios e feedback (Deterding et al., 2011). Sínteses de pesquisa relatam efeitos variados conforme a intervenção, o público e o desfecho considerado; não se deve transferir uma média de estudos para uma aplicação específica sem avaliação própria (Bai; Hew; Huang, 2020). No LixoCerto, quiz, pontuação, selos e desafios estruturam interações e feedbacks. A avaliação técnica verifica se esses estados aparecem e persistem nos fluxos testados, não se produzem engajamento duradouro.')

    add_heading(doc, '2.3 Usabilidade por inspeção, acessibilidade e desempenho web', 2)
    add_text_paragraph(doc,
        'As dez heurísticas de Nielsen oferecem princípios de inspeção para identificar problemas de visibilidade do estado, correspondência com o mundo real, controle, consistência, prevenção de erros, reconhecimento, flexibilidade, minimalismo, recuperação de erros e ajuda (Nielsen, 1994). Neste trabalho atribuímos severidade de 0 a 4, para priorizar achados de inspeção; essa classificação não mede frequência, impacto observado em participantes ou satisfação percebida.')
    add_text_paragraph(doc,
        'A WCAG 2.2 estabelece critérios de sucesso para conteúdo web acessível, organizados em níveis A, AA e AAA (W3C, 2023). Ferramentas automatizadas detectam apenas parte dos problemas: navegação por teclado, uso de leitor de tela, compreensão, foco e diversas decisões de interação exigem inspeção manual e, idealmente, avaliação com pessoas. Assim, relatamos separadamente violações confirmadas e regras incompletas da análise automática.')
    add_text_paragraph(doc,
        'Para Core Web Vitals, os limites de referência de bom desempenho são LCP até 2,5 s, INP até 200 ms e CLS até 0,1, avaliados em dados reais no percentil 75 quando disponíveis (Google, s.d.). Lighthouse permite uma avaliação de laboratório reproduzível, mas métricas simuladas não substituem dados de campo. A API Geolocation descreve `accuracy` em metros como raio de confiança de 95% (W3C, 2024); por isso, a aceitação de precisão reportada de até 10 m combinada a um raio operacional de 5 m demanda interpretação e ensaios físicos.')

    add_heading(doc, '2.4 Trabalhos correlatos', 2)
    add_text_paragraph(doc,
        'Realizamos comparação documental das funções apresentadas em páginas oficiais de serviços relacionados. A Tabela 1 delimita sobreposições funcionais; não comparamos qualidade, desempenho, cobertura real, popularidade, eficácia educativa ou resultados ambientais.')
    add_table(doc, 1, 'Comparação funcional documental de soluções relacionadas',
        ['Solução', 'Funções descritas nas fontes consultadas', 'Relação com o LixoCerto', 'Limite da observação'], [
            ['Cataki', 'Conecta geradores a profissionais da reciclagem; descreve separação, localização, combinação de coleta e remuneração.', 'Sobreposição parcial em encaminhamento e orientação para destinação.', 'Não avaliamos operação, cobertura ou experiência do serviço.'],
            ['eCycle', 'Localizador de postos por material e localidade, com busca por endereço/CEP ou localização.', 'Sobreposição na busca de pontos e categorias.', 'A consulta documental não verificou completude nem atualização dos resultados locais.'],
            ['Recicla Sampa', 'Página de pontos com mapa, busca e visualização em lista para a cidade de São Paulo.', 'Sobreposição na descoberta geográfica de locais.', 'A sessão mostrou interface de mapa, mas não carregou marcadores; não inferimos cobertura.'],
            ['LixoCerto', 'Guia com instruções e fontes, pontos, quiz, desafios, progresso e check-in com localização e foto.', 'Integra orientação, atividade gamificada e registro de descarte.', 'Neste estudo verificamos execução técnica dos fluxos, não impacto nem eficácia comparativa.'],
        ], [2.2, 5.4, 4.2, 4.2], 'páginas oficiais de Cataki, eCycle, Recicla Sampa e LixoCerto consultadas em 30 out. 2026', 8.2)
    add_text_paragraph(doc,
        'A comparação situa a contribuição do LixoCerto na integração, em um mesmo artefato, de orientação por material, descoberta de pontos e tarefas gamificadas com registro de descarte. As fontes consultadas descrevem funções, mas não sustentam superioridade ou eficácia comparativa; por isso, avaliamos a execução das jornadas próprias em vez de ranquear os serviços.')

    add_heading(doc, '3 MATERIAIS E MÉTODOS', 1)
    add_heading(doc, '3.1 Delineamento e desenho de pesquisa', 2)
    add_text_paragraph(doc,
        'Conduzimos pesquisa aplicada de desenvolvimento e avaliação de artefato, orientada pelo paradigma Design Science Research (DSR), que vincula problema relevante, construção e avaliação de soluções de tecnologia da informação (Hevner et al., 2004; Dresch; Lacerda; Antunes Júnior, 2015). Operacionalizamos o estudo em cinco movimentos: delimitação do problema e dos requisitos; definição da solução; desenvolvimento incremental; demonstração em tarefas; e avaliação técnica e comunicação dos resultados.')
    add_figure(doc, 1, EVIDENCE/'diagramas/desenho-pesquisa-dsr.png',
        'Desenho de pesquisa orientado por Design Science Research',
        'elaboração própria com base em Hevner et al. (2004) e Dresch, Lacerda e Antunes Júnior (2015)',
        'Fluxo da pesquisa: problema e requisitos, projeto e desenvolvimento do artefato, demonstração em tarefas, avaliação técnica e comunicação.', 5.7)
    add_text_paragraph(doc,
        'A Figura 1 relaciona construção e avaliação do artefato sem pressupor validação com participantes. A demonstração utiliza contas e dados de teste sintéticos; os resultados respondem se as jornadas previstas puderam ser executadas e quais achados técnicos permaneceram na versão avaliada.')

    add_heading(doc, '3.2 Público-alvo, requisitos e processo de desenvolvimento', 2)
    add_text_paragraph(doc,
        'Definimos como público pretendido pessoas que buscam orientação para separar resíduos e localizar pontos de recebimento, com administradores responsáveis pela curadoria do guia e revisão de registros. Os requisitos operacionais contemplam consulta por material e localidade, apresentação de instruções e fontes, quiz e pontuação, desafios e registro de descarte com validação de localização e fotografia. Esses requisitos refletem o escopo do artefato e das tarefas avaliadas, não entrevistas ou observação de usuários.')
    add_text_paragraph(doc,
        'Desenvolvemos o sistema em incrementos por módulos, com controle de versão, contratos entre serviços, testes automatizados e ciclos de correção. Não atribuímos o processo a Scrum ou Kanban formal porque os artefatos examinados não permitem afirmar a existência de sprints, cerimônias ou quadro de fluxo. Modelamos os atores e as funcionalidades principais no diagrama de casos de uso e as entidades centrais pelas relações documentadas no banco.')
    add_figure(doc, 2, EVIDENCE/'diagramas/arquitetura-logica.png',
        'Arquitetura lógica da plataforma LixoCerto',
        'elaboração própria a partir do código e da infraestrutura dos repositórios',
        'Arquitetura com interface pública React/Vite, Envoy gRPC-Web e API NestJS como núcleo das regras; os módulos internos e integrações de dados não são acessados diretamente pelo navegador.', 6.0)
    add_text_paragraph(doc,
        'A arquitetura adota a API como núcleo das regras de negócio: o navegador não acessa diretamente os módulos internos nem as integrações de persistência. React e Vite estruturam a interface pública em componentes e dividem o código por rotas; contratos gRPC tipados e Envoy com gRPC-Web intermediam chamadas aos serviços NestJS. PostgreSQL mantém as relações persistentes, Redis atende funções de cache e o armazenamento compatível com S3 guarda evidências binárias. Essas tecnologias correspondem a responsabilidades distintas do artefato e não implicam que todos os serviços externos estejam implantados em produção.')
    add_figure(doc, 3, EVIDENCE/'diagramas/casos-de-uso.png',
        'Casos de uso das pessoas usuárias e da administração',
        'elaboração própria com base nos módulos e permissões observados',
        'Diagrama de casos de uso agrupando consulta ao guia e pontos, quiz, desafios e check-in para a pessoa usuária; e curadoria, moderação e gestão administrativa.', 5.7)
    add_figure(doc, 4, EVIDENCE/'diagramas/modelo-er-core.png',
        'Modelo ER conceitual dos módulos centrais',
        'elaboração própria; relações reproduzidas a partir das chaves estrangeiras observadas nas migrations',
        'Modelo entidade-relacionamento com USUARIO, MATERIAL, PONTO_COLETA, QUESTAO, ALTERNATIVA, QUIZ_ATRIBUICAO, TENTATIVA_QUIZ, CHECKIN_DESCARTE e REGISTRO_DESCARTE; setas indicam as cardinalidades observadas.', 5.8)
    add_text_paragraph(doc,
        'As Figuras 3 e 4 sintetizam os papéis e relações persistidas considerados no recorte. O ER é intencionalmente parcial: representa entidades e chaves estrangeiras centrais à demonstração e não substitui o esquema completo do banco.')

    add_heading(doc, '3.3 Instrumentos e protocolo de avaliação', 2)
    add_text_paragraph(doc,
        'Executamos os ensaios no sandbox local, com PostgreSQL e Redis descartáveis, API NestJS, Envoy gRPC-Web e build do frontend. A conta de avaliação, a foto PNG e o ponto EMLUR reproduzido no banco eram fixtures; as coordenadas foram simuladas no navegador. O provedor de e-mail foi configurado em modo de registro local, sem envio a destinatários reais. Os testes não acessaram a VM nem a homologação.')
    add_text_paragraph(doc,
        'A suíte Playwright verificou estados visíveis e transições de cadastro/login, onboarding, busca no guia, busca de pontos, quiz e pontuação, adesão a desafio, progresso, geofence e envio de foto. Executamos também os 11 testes de integração existentes e o E2E de saúde gRPC da API. Para Axe, analisamos sete telas com tags WCAG 2.2 A/AA e boas práticas em viewport desktop 1280×900 CSS px. Para Lighthouse User Flow, repetimos três vezes navegação e interações na build de produção em emulação móvel 390×844 CSS px, DPR 1, RTT 150 ms, download 1.474,56 Kbps, upload 675 Kbps e CPU ×4.')
    add_text_paragraph(doc,
        'Realizamos dez inspeções heurísticas independentes, uma por heurística de Nielsen. Classificamos severidade em escala de 0 a 4: 0, sem problema; 1, cosmético; 2, menor; 3, maior; e 4, catastrófico. Registramos separadamente se cada achado permanecia presente ou havia sido parcialmente mitigado. Por fim, comparamos o carregamento de `/login` em duas builds equivalentes: alternamos apenas `App.tsx` entre a versão de referência sem lazy loading e a versão com divisão de rotas, medindo os arquivos JavaScript realmente solicitados e a transferência gzip.')
    add_text_paragraph(doc,
        'Para reprodução, registramos os SHAs de referência, versões das ferramentas, comandos, configuração de emulação e limites do ensaio no suplemento técnico do repositório (`reports/etapa2-evidencias/rastreabilidade.md`), juntamente com os scripts e relatórios brutos. A comparação de bundle usa uma base anterior identificada e o estado avaliado, com alterações restritas ao carregamento de rotas.')

    add_heading(doc, '4 RESULTADOS E DISCUSSÃO', 1)
    add_heading(doc, '4.1 Módulos implementados e decisões de projeto', 2)
    add_text_paragraph(doc,
        'Construímos cinco módulos centrais: Guia de descarte com categorias, instruções, cuidados e fontes; busca de pontos por localidade, categorias aceitas e localização; quiz com resposta e pontuação; desafios e painel de progresso; e check-in com confirmação de presença, fotografia por categoria e encaminhamento para revisão. Além desses fluxos, a plataforma contempla autenticação, cadastro e perfil, gestão administrativa de usuários e operações CRUD para os cadastros do sistema.')
    add_text_paragraph(doc,
        'Também implementamos um roteiro de primeiro acesso que conduz a pessoa por missões guiadas: iniciar o percurso, buscar um material no Guia e consultar um ponto de coleta, liberando a navegação ao concluir as etapas. A jornada de cadastro, login e onboarding verificou a sequência e o acesso às áreas; não a tratamos como evidência de aprendizagem nem de melhoria percebida por participantes.')
    add_text_paragraph(doc,
        'No check-in, a implementação estabelece a expiração de fotos e coordenadas 90 dias após a decisão de revisão, desde que não haja outro registro pendente associado ao mesmo check-in. Um processo do serviço inicia o expurgo ao subir e volta a executá-lo a cada dez minutos; quando vence a retenção, remove a imagem do armazenamento, limpa a chave no registro e elimina coordenadas. Essa é uma decisão de projeto codificada, não um resultado de teste de espera por 90 dias ou de restauração de objetos.')

    add_heading(doc, '4.2 Testes ponta a ponta e testes da API', 2)
    add_text_paragraph(doc,
        'A Tabela 2 registra as sete jornadas executadas. Todas as asserções foram verificadas contra estados da interface e chamadas da API no ambiente local descartável; os dados de conta, localização e imagem eram fictícios.')
    add_table(doc, 2, 'Jornadas Playwright e resultados de execução',
        ['Jornada', 'Passos e checkpoints observados', 'Resultado'], [
            ['Cadastro, login e onboarding', 'Criação com e-mail .test; tela de confirmação sem envio real; login fictício; missões guiadas para buscar material e ponto.', 'Aprovado (1/1)'],
            ['Compatibilidade de login', 'Autenticação de conta fictícia legada com senha válida anterior à política atual de cadastro.', 'Aprovado (1/1)'],
            ['Consulta ao Guia', 'Busca de papel/papelão; abertura da ficha; instruções, cuidados e links de fonte.', 'Aprovado (1/1)'],
            ['Busca de pontos', 'Consulta por João Pessoa; estado Verificado; categoria, atendimento, fonte e ressalva do pin.', 'Aprovado (1/1)'],
            ['Quiz, desafio e progresso', 'Resposta correta, feedback de +20 pontos; adesão a desafio; atualização do painel com pontos e desafio ativo.', 'Aprovado (1/1)'],
            ['Geofence fora do raio', 'Coordenada de navegador simulada a aproximadamente 23 m do pin; API rejeita limite de 5 m e apresenta instrução de aproximação.', 'Aprovado (1/1)'],
            ['Check-in com foto', 'Coordenada simulada no pin, confirmação no raio, anexo de PNG sintético por categoria e envio com estado Em revisão.', 'Aprovado (1/1)'],
            ['Total Playwright', 'Sete testes automatizados, sem casos ignorados, falhos ou flaky na execução registrada.', '7/7 aprovados'],
        ], [3.0, 10.2, 2.8], 'relatório Playwright arquivado em `raw/playwright.json`; execução local', 8.4)
    add_text_paragraph(doc,
        'O roteiro de primeiro acesso integra essa evidência funcional: depois do cadastro e login, a interface apresenta tarefas guiadas de consulta ao Guia e aos pontos de coleta antes de liberar o uso geral. Assim, registramos a implementação e a execução técnica do onboarding como parte do resultado, sem afirmar que ele melhora a orientação ou a retenção de usuários.')
    add_text_paragraph(doc,
        'Na mesma configuração local, executamos 11/11 testes de integração da API e 1/1 teste E2E de saúde gRPC. As verificações complementares passaram: 122 testes unitários da API em 33 arquivos, 17 testes unitários do frontend em três arquivos, análise de tipos, validação dos contratos Buf e builds de produção. Esses resultados cobrem os comandos e dados descritos, não todas as combinações de entrada, falhas externas ou ambientes de implantação.')

    add_heading(doc, '4.3 Avaliação heurística de Nielsen', 2)
    add_text_paragraph(doc,
        'A inspeção pós-correções identificou 32 achados nas dez heurísticas: um de severidade 1, 28 de severidade 2 e três de severidade 3; nenhum recebeu severidade 4. Dezesseis achados permaneceram presentes e dezesseis foram parcialmente mitigados, sem que a mitigação parcial represente resolução. A Tabela 3 resume o número por heurística e os principais temas.')
    heuristic_rows = [
        ['1. Visibilidade do estado', '0 / 3 / 0 / 0', '2 presentes; 1 parcial', 'Pendência de ações administrativas e distinção entre falha de mutação e falha de atualização.'],
        ['2. Correspondência com o mundo real', '0 / 2 / 1 / 0', '1 presente; 2 parciais', 'Código de cores de resíduos perigosos; ressalva local e significado do estado Verificado.'],
        ['3. Controle e liberdade', '0 / 6 / 0 / 0', '3 presentes; 3 parciais', 'Retomar onboarding, sair de missão e recuperar formulário/GPS após cancelamento ou erro.'],
        ['4. Consistência e padrões', '1 / 2 / 0 / 0', '2 presentes; 1 parcial', 'Diferenças entre filtros e feedbacks; inconsistência de reticências.'],
        ['5. Prevenção de erros', '0 / 7 / 0 / 0', '5 presentes; 2 parciais', 'Estado de ponto obsoleto, validações de check-in, repetição de submissão e divergência de arquivos.'],
        ['6. Reconhecimento', '0 / 1 / 0 / 0', '0 presentes; 1 parcial', 'O feedback do quiz não preserva enunciado e alternativas que geraram o resultado.'],
        ['7. Flexibilidade e eficiência', '0 / 2 / 0 / 0', '1 presente; 1 parcial', 'Sugestões de busca preenchem o campo sem iniciar a consulta; adiamento não persiste.'],
        ['8. Estética e minimalismo', '0 / 2 / 0 / 0', '0 presentes; 2 parciais', 'Listas com área vazia e instruções extensas que deslocam ações.'],
        ['9. Diagnóstico e recuperação', '0 / 1 / 2 / 0', '1 presente; 2 parciais', 'Confirmação de resposta, falha de recarga ambígua e ausência de retry/indicador de dados antigos.'],
        ['10. Ajuda e documentação', '0 / 2 / 0 / 0', '1 presente; 1 parcial', 'Retomada prometida sem controle claro; ajuda geral pouco contextual.'],
        ['Total (S1 / S2 / S3 / S4)', '1 / 28 / 3 / 0', '16 presentes; 16 parciais', '32 achados; escala de gravidade 0–4.'],
    ]
    add_table(doc, 3, 'Achados por heurística, severidade e situação',
        ['Heurística', 'Sev. 1 / 2 / 3 / 4', 'Situação', 'Evidência principal'], heuristic_rows,
        [3.0, 2.4, 3.0, 7.6], 'dez inspeções independentes documentadas em `avaliacao-heuristica-etapa2.md`', 7.9)
    add_text_paragraph(doc,
        'Os três achados de severidade 3 foram: uso de roxo para resíduos perigosos, apesar da convenção nacional de cores associar essa cor a resíduos radioativos e reservar laranja para perigosos (Brasil, 2001); possibilidade de a interface apresentar uma resposta de quiz como registrada sem confirmação semântica da API; e mensagem de erro após mutação aceita pelo servidor quando a recarga subsequente falha. A gravidade é prioridade inferida por inspeção, não taxa de ocorrência com participantes.')

    add_heading(doc, '4.4 Acessibilidade automática', 2)
    add_text_paragraph(doc,
        'A análise Axe abrangeu login, confirmação de e-mail, visão geral, Guia, pontos, quiz e check-in. A Tabela 4 mostra os resultados por tela. Os 283 resultados de regra aprovados são agregados das sete execuções; seis regras `color-contrast` ficaram incompletas e exigem decisão manual. Nenhuma tela apresentou violação confirmada nesta passagem.')
    axe_rows = [
        ['Login', '0', '42', '1 — color-contrast'],
        ['Confirmação de e-mail', '0', '24', '0'],
        ['Visão geral', '0', '44', '1 — color-contrast'],
        ['Guia', '0', '44', '1 — color-contrast'],
        ['Pontos', '0', '45', '1 — color-contrast'],
        ['Quiz', '0', '39', '1 — color-contrast'],
        ['Check-in', '0', '45', '1 — color-contrast'],
        ['Total', '0', '283', '6 regras incompletas'],
    ]
    add_table(doc, 4, 'Varredura Axe nas telas principais',
        ['Tela', 'Violações confirmadas', 'Resultados de regra aprovados', 'Regras incompletas'], axe_rows,
        [4.0, 3.0, 4.0, 5.0], 'relatório Axe com tags WCAG 2.2 A/AA e boas práticas, em `raw/axe-wcag22.json`', 8.5)
    add_text_paragraph(doc,
        'O resultado zero se restringe ao que a ferramenta conseguiu avaliar automaticamente nas telas e estados capturados. A incompletude de contraste não é aprovação: gradientes, conteúdo dinâmico e combinações específicas ainda requerem inspeção manual. Também não realizamos nesta rodada auditoria completa com teclado, leitor de tela, ampliação, alto contraste ou tecnologia assistiva.')

    add_heading(doc, '4.5 Desempenho e code splitting', 2)
    add_text_paragraph(doc,
        'Repetimos três vezes a avaliação Lighthouse User Flow em build de produção sob emulação móvel. A Tabela 5 apresenta as medianas de score Lighthouse e das métricas obtidas em navegação/interação; o valor de LCP corresponde à navegação inicial de login, enquanto INP é reportado por timespan de interação.')
    lighthouse_rows = [
        ['Navegação `/login`', '94 (94–100)', 'LCP 2.930,6 ms', 'CLS 0,000'],
        ['Autenticação e primeiro acesso', '69 (68–69)', 'INP 58,1 ms', '—'],
        ['Busca no Guia', '97 (97–97)', 'INP 230,5 ms', '—'],
        ['Busca de ponto', '97 (96–97)', 'INP 245,7 ms', '—'],
        ['Abertura do quiz', '100 (100–100)', 'INP 158,7 ms', '—'],
        ['Início do check-in', '98 (98–99)', 'INP 202,6 ms', '—'],
    ]
    add_table(doc, 5, 'Métricas Lighthouse em três execuções de laboratório',
        ['Etapa', 'Score mediano (intervalo)', 'Métrica e mediana', 'CLS'], lighthouse_rows,
        [4.3, 3.7, 4.5, 2.5], 'relatórios `raw/lighthouse-run-1.json` a `raw/lighthouse-run-3.json`; emulação 390×844 CSS px, DPR 1, RTT 150 ms, download 1.474,56 Kbps, upload 675 Kbps, CPU ×4', 8.1)
    add_text_paragraph(doc,
        'O LCP mediano foi 2,93 s: acima do limite de referência de bom desempenho (2,5 s), embora as três amostras tenham ficado abaixo de 3 s. A mediana excede 200 ms em três das cinco interações medidas (busca no Guia, busca de ponto e início do check-in); o CLS foi zero. As pontuações e métricas são de laboratório e não representam o percentil 75 de sessões reais. Assim, o ensaio atende apenas a um requisito de latência observado sob a configuração simulada; não permite afirmar aprovação em Core Web Vitals de campo.')
    add_text_paragraph(doc,
        'Na comparação controlada do code splitting, a transferência inicial de JavaScript em `/login` caiu de 797.683 para 474.077 bytes sem gzip e de 226.801 para 146.583 bytes com gzip. A redução foi de 40,6% sem gzip e 35,4% em gzip, com mudança de um para cinco arquivos solicitados. A build final soma 800.223 bytes de JavaScript distribuídos entre chunks; a melhoria incide sobre a transferência inicial da rota e adia código das demais áreas.')
    add_table(doc, 6, 'Comparação do carregamento inicial da rota `/login`',
        ['Medida', 'Antes do lazy loading', 'Após code splitting', 'Variação'], [
            ['JavaScript solicitado, sem gzip', '797.683 bytes', '474.077 bytes', '−40,6%'],
            ['Transferência dos scripts em gzip', '226.801 bytes', '146.583 bytes', '−35,4%'],
            ['Arquivos JavaScript solicitados', '1', '5', 'chunks compartilhados e da rota'],
            ['JavaScript na build completa (após)', '—', '800.223 bytes', 'redistribuído em chunks; não removido'],
        ], [5.2, 3.7, 3.7, 3.4], 'comparação controlada do mesmo working tree, alternando somente o `App.tsx`; dados em `raw/bundle-comparison.json`', 8.5)

    add_heading(doc, '4.6 Jornadas observadas nas telas', 2)
    add_text_paragraph(doc,
        'As Figuras 5 a 9 apresentam telas na ordem da jornada. As Figuras 5 e 6 são recortes em tons de cinza da homologação, obtidos em consulta de leitura e higienizados para excluir identidade administrativa e registros alheios; as Figuras 7 a 9 provêm de execuções Playwright locais com conta fictícia e dados sintéticos. Cada figura documenta somente o estado visível indicado, não cobertura de todas as funções.')
    add_figure(doc, 5, EVIDENCE/'figuras/figura-03-guia.png',
        'Busca e ficha publicada no Guia de descarte',
        'consulta somente de leitura ao Guia de homologação (LixoCerto, 2026); recortes sem identidade administrativa e composição em escala de cinza',
        'Capturas lado a lado da busca por papelão e da ficha publicada Papel e papelão secos para coleta seletiva, com instruções de descarte, preparação e cuidados; identidade administrativa excluída.', 6.0)
    add_text_paragraph(doc,
        'A Figura 5 documenta, na homologação, uma busca por “papelão” que retornou a ficha publicada “Papel e papelão secos para coleta seletiva”. O recorte exclui a identidade da conta administrativa e nenhum conteúdo foi alterado nessa consulta. A ficha apresenta preparo, cuidados e encaminhamento condicionado à regra municipal ou do operador; as fontes estão identificadas no registro publicado. No conjunto do sistema, a heurística de correspondência com o mundo real reuniu três achados; nesta tela, uma orientação municipal pode ficar pouco saliente. A captura não demonstra atualização contínua das regras locais.')
    add_figure(doc, 6, EVIDENCE/'figuras/figura-04-pontos.png',
        'Mapa do pin cartográfico e cartão EMLUR exibido na homologação',
        'mapa © OpenStreetMap contributors (ODbL); cartão recortado da homologação em 30 out. 2026; coordenada de referência obtida no Google Maps',
        'Mapa em tons de cinza centrado na EMLUR e recorte do cartão homologado com estado Verificado, endereço, horário listado, categoria de eletroeletrônicos e informações do local.', 6.0)
    add_text_paragraph(doc,
        'Na Figura 6, a Prefeitura de João Pessoa sustenta o endereço e o recebimento de resíduos eletroeletrônicos (Prefeitura de João Pessoa, 2026); o horário de segunda a sexta, 8h–14h, foi consultado no Waze e precisa ser confirmado antes da visita (Waze, s.d.). O cartão foi aprovado como “Verificado” na homologação, mas o pin deriva de uma ficha cartográfica do Google Maps e não foi medido na entrada (Google Maps, s.d.). Portanto, o status cadastral não certifica precisão geodésica, operação atual nem aceitação incondicional. O OpenStreetMap foi usado para a visualização cartográfica e é creditado na legenda.')
    add_figure(doc, 7, EVIDENCE/'figuras/figura-05-quiz.png',
        'Questão do quiz e feedback de resposta com pontuação',
        'capturas da jornada Playwright local com questão publicada e conta fictícia',
        'Telas lado a lado: pergunta sobre descarte de pilhas e resultado com resposta correta, explicação e mais vinte pontos.', 6.0)
    add_text_paragraph(doc,
        'A Figura 7 evidencia feedback de acerto e pontuação na tarefa executada. A inspeção identificou um achado de severidade 2 na heurística de reconhecimento: o feedback não preserva o enunciado e as alternativas que originaram a resposta, reduzindo a possibilidade de revisão posterior. A imagem comprova o caminho de acerto da fixture, não retenção do conhecimento.')
    add_figure(doc, 8, EVIDENCE/'figuras/figura-06-progresso.png',
        'Painel de progresso com pontos, nível, selos e desafio ativo',
        'captura Playwright local com conta fictícia após quiz e participação em desafio',
        'Painel sintético exibindo 20 pontos, nível Iniciante, sequência de um dia, zero de cinco selos e barra de progresso de um desafio ativo.', 6.0)
    add_text_paragraph(doc,
        'Na Figura 8, a interface reúne pontos, nível, sequência, selos e desafio ativo em uma tela, após a ação do quiz e a adesão ao desafio. Os valores pertencem a uma conta de teste. A atualização do painel foi verificada pela jornada ponta a ponta; a avaliação heurística, contudo, encontrou questões de visibilidade de estado e retomada de onboarding descritas na Tabela 3.')
    add_figure(doc, 9, EVIDENCE/'figuras/figura-07-checkin.png',
        'Check-in móvel com rejeição fora do raio e anexação de foto sintética',
        'capturas Playwright em viewport móvel 390×844 CSS px; coordenadas simuladas e PNG artificial',
        'Duas telas móveis: mensagem acionável de rejeição a cerca de 23 metros quando o limite é 5 metros; formulário com foto sintética anexada e aviso de retenção após revisão.', 6.0)
    add_text_paragraph(doc,
        'A Figura 9 reúne a rejeição simulada fora do raio e o envio de fotografia sem pessoas, placas ou dados reais. O sistema informou a distância e orientou a aproximação, e o cenário dentro do raio permitiu enviar o registro para revisão. Como a especificação de localização descreve a precisão reportada como raio de confiança de 95%, aceitar precisão de até 10 m em uma geofence de 5 m pode bloquear uma pessoa que esteja fisicamente próxima ao ponto (W3C, 2024). Esse risco é operacional e não pode ser quantificado pelo ensaio com coordenadas simuladas.')

    add_heading(doc, '4.7 Publicação de conteúdo e discussão integrada', 2)
    add_text_paragraph(doc,
        'Na primeira observação da homologação, o Guia mostrava oito conteúdos; em reconsulta de leitura em 30/10/2026, exibiu nove (LixoCerto, 2026). A busca por “papelão” retornou uma ficha publicada, conforme a Figura 5. Os cinco temas revisados documentalmente — latas, papel/papelão, embalagens plásticas, compostagem e vidro — já constavam da lista; não alteramos status nem duplicamos fichas. As revisões textuais e um seeder idempotente foram preparados localmente, preservando o status de aprovação, mas não foram aplicados remotamente. A diferença entre oito e nove é tratada como variação entre consultas, não como publicação feita por nós. O valor de 54 materiais inicialmente em rascunho pertence ao seeder de desenvolvimento e não deve ser comparado diretamente ao contador de homologação.')
    add_text_paragraph(doc,
        'Em conjunto, os resultados mostram que as principais jornadas especificadas foram executadas na versão local, que a API passou nos testes de integração e saúde existentes e que o code splitting reduziu a transferência inicial de JavaScript. A inspeção heurística mantém problemas de prevenção de erro, controle, clareza semântica e recuperação, mesmo após melhorias de interface; a varredura automatizada não substitui avaliação manual. O LCP mediano próximo de 3 s e três interações acima de 200 ms indicam prioridades de desempenho, enquanto o conflito entre raio de 5 m e acurácia reportada de até 10 m recomenda teste com aparelhos e posições reais antes de ajustar a regra. Nenhum desses resultados, isoladamente, demonstra eficácia ambiental, aprendizagem ou melhora percebida de usabilidade.')

    add_heading(doc, '4.8 Limitações da avaliação', 2)
    add_text_paragraph(doc,
        'A avaliação não incluiu participantes, observação de tarefas, aplicação do SUS ou medida de aprendizagem; por isso, não estimamos satisfação, taxa de sucesso de usuários, mudança de comportamento ou eficácia ambiental. O ensaio de geofence foi feito com coordenadas do navegador, não com celulares em campo, interiores ou pontos diferentes; a coordenada EMLUR é cartográfica, não aferida na entrada. O Axe avaliou sete telas por regras automáticas, sem auditoria completa de teclado, leitor de tela ou ampliação. Lighthouse forneceu laboratório sintético, não dados CrUX no percentil 75. A política de expurgo de 90 dias foi verificada no código, mas não aguardamos sua execução real nem testamos recuperação de imagens e coordenadas expiradas. A comparação de catálogo foi documental e não verificou a operação diária dos pontos.')

    add_heading(doc, '5 CONSIDERAÇÕES FINAIS', 1)
    add_text_paragraph(doc,
        'Desenvolvemos e avaliamos tecnicamente a plataforma LixoCerto, reunindo guia de descarte, busca de pontos, quiz, desafios e check-in com fotografia. Na configuração local, aprovamos sete jornadas Playwright, 11 testes de integração e um E2E gRPC; identificamos 32 achados heurísticos, nenhuma violação automática confirmada pelo Axe nas sete telas e seis regras de contraste pendentes de julgamento manual. Lighthouse registrou LCP mediano de 2,93 s, CLS zero e medianas de INP de 58 a 246 ms; o code splitting diminuiu em 35,4% a transferência gzip inicial de JavaScript em `/login`. A consulta mais recente à homologação mostrou nove fichas publicadas (após uma observação anterior de oito) e confirmou o cadastro autorizado da EMLUR como verificado, com pin explicitamente tratado como cartográfico e não medido em campo.')
    add_text_paragraph(doc,
        'Essas evidências sustentam o funcionamento dos fluxos e decisões técnicas na versão e no ambiente descritos, sem concluir que o sistema melhora descarte, aprendizagem ou usabilidade percebida. Como trabalhos futuros, priorizamos: (1) avaliar as tarefas com participantes do público-alvo, após avaliação ética e consentimento, incluindo instrumentos de usabilidade e aprendizagem; (2) testar geofence em aparelhos e locais reais, registrando acurácia e rejeições; (3) complementar a acessibilidade com teclado, leitor de tela e ampliação; e (4) testar em homologação a retenção, o expurgo e a restauração conjunta de banco e objetos.')

    add_heading(doc, 'REFERÊNCIAS', 1)
    references = [
        ('ABREMA — ASSOCIAÇÃO BRASILEIRA DE RESÍDUOS E MEIO AMBIENTE. ', 'Panorama dos resíduos sólidos no Brasil 2025: ano-base 2024', '. São Paulo: ABREMA, 2025.'),
        ('BAI, S.; HEW, K. F.; HUANG, B. ', 'Does gamification improve student learning outcome? Evidence from a meta-analysis and synthesis of qualitative data in educational contexts', '. Educational Research Review, v. 30, art. 100322, 2020. DOI: 10.1016/j.edurev.2020.100322.'),
        ('BRASIL. ', 'Lei nº 9.795, de 27 de abril de 1999', '. Dispõe sobre a educação ambiental e institui a Política Nacional de Educação Ambiental. Brasília, DF, 1999.'),
        ('BRASIL. ', 'Lei nº 12.305, de 2 de agosto de 2010', '. Institui a Política Nacional de Resíduos Sólidos. Brasília, DF, 2010.'),
        ('BRASIL. Conselho Nacional do Meio Ambiente. ', 'Resolução nº 275, de 25 de abril de 2001', '. Estabelece código de cores para os diferentes tipos de resíduos na coleta seletiva. Brasília, DF, 2001.'),
        ('BRASIL. ', 'Decreto nº 10.240, de 12 de fevereiro de 2020', '. Regulamenta sistema de logística reversa de produtos eletroeletrônicos e seus componentes. Brasília, DF, 2020.'),
        ('BRASIL. ', 'Lei nº 13.709, de 14 de agosto de 2018', '. Lei Geral de Proteção de Dados Pessoais. Brasília, DF, 2018.'),
        ('CATaki. ', 'Como funciona', '. Disponível em: https://www.cataki.org/. Acesso em: 30 out. 2026.'),
        ('DETERDING, S. et al. ', 'From game design elements to gamefulness: defining “gamification”', '. In: INTERNATIONAL ACADEMIC MINDTREK CONFERENCE, 15., 2011, Tampere. Proceedings [...]. New York: ACM, 2011. p. 9–15. DOI: 10.1145/2181037.2181040.'),
        ('DRESCH, A.; LACERDA, D. P.; ANTUNES JÚNIOR, J. A. V. ', 'Design Science Research: método de pesquisa para avanço da ciência e tecnologia', '. Porto Alegre: Bookman, 2015.'),
        ('eCYCLE. ', 'Postos de reciclagem', '. Disponível em: https://www.ecycle.com.br/postos/reciclagem.php. Acesso em: 30 out. 2026.'),
        ('GIL, A. C. ', 'Como elaborar projetos de pesquisa', '. 7. ed. Barueri: Atlas, 2022.'),
        ('GOOGLE. ', 'Web Vitals', '. web.dev. Disponível em: https://web.dev/articles/vitals. Acesso em: 30 out. 2026.'),
        ('GOOGLE MAPS. ', 'EMLUR — Municipal Urban Cleaning Company, João Pessoa [mapa]', '. Disponível em: https://www.google.com/maps/search/?api=1&query=EMLUR%2C%20177%20Avenida%20Minas%20Gerais%2C%20Jo%C3%A3o%20Pessoa%2C%20PB. Acesso em: 30 out. 2026.'),
        ('HAJJ-HASSAN, M.; CHAKER, R.; CEDERQVIST, A. ', 'Environmental education: a systematic review on the use of digital tools for fostering sustainability awareness', '. Sustainability, v. 16, n. 9, 2024.'),
        ('HEVNER, A. R.; MARCH, S. T.; PARK, J.; RAM, S. ', 'Design science in information systems research', '. MIS Quarterly, v. 28, n. 1, p. 75–105, 2004. DOI: 10.2307/25148625.'),
        ('ITU — INTERNATIONAL TELECOMMUNICATION UNION; UNITAR — UNITED NATIONS INSTITUTE FOR TRAINING AND RESEARCH. ', 'Brazil: 2022 e-waste statistics', '. Global E-waste Statistics Partnership, 2024. Disponível em: https://globalewaste.org/statistics/country/brazil/2022/. Acesso em: 30 out. 2026.'),
        ('LIXOCERTO. ', 'Guia de descarte', '. Plataforma web em homologação. Disponível em: https://centrallixocerto.com.br/. Acesso em: 30 out. 2026.'),
        ('NIELSEN, J. ', 'Enhancing the explanatory power of usability heuristics', '. In: CONFERENCE ON HUMAN FACTORS IN COMPUTING SYSTEMS, 1994, Boston. Proceedings [...]. New York: ACM, 1994. p. 152–158. DOI: 10.1145/191666.191729.'),
        ('OPENSTREETMAP CONTRIBUTORS. ', 'OpenStreetMap', '. Dados cartográficos disponibilizados sob licença ODbL. Disponível em: https://www.openstreetmap.org/. Acesso em: 30 out. 2026.'),
        ('PREFEITURA DE JOÃO PESSOA. ', 'Recicla JP dá novo destino a resíduos eletrônicos e amplia ações de inclusão social', '. João Pessoa, 29 set. 2026. Disponível em: https://joaopessoa.pb.gov.br/noticias/recicla-jp-da-novo-destino-a-residuos-eletronicos-e-amplia-acoes-de-inclusao-social. Acesso em: 30 out. 2026.'),
        ('RECICLA SAMPA. ', 'Pontos de coleta', '. Disponível em: https://www.reciclasampa.com.br/pontos-de-coleta. Acesso em: 30 out. 2026.'),
        ('W3C — WORLD WIDE WEB CONSORTIUM. ', 'Web Content Accessibility Guidelines (WCAG) 2.2', '. W3C Recommendation, 5 out. 2023. Disponível em: https://www.w3.org/TR/WCAG22/. Acesso em: 30 out. 2026.'),
        ('W3C — WORLD WIDE WEB CONSORTIUM. ', 'Geolocation', '. W3C Recommendation, 2024. Disponível em: https://www.w3.org/TR/geolocation/. Acesso em: 30 out. 2026.'),
        ('WAZE. ', 'EMLUR — Avenida Minas Gerais, 177, João Pessoa', '. Diretório cartográfico. Disponível em: https://www.waze.com/live-map/directions/emlur-av.-minas-gerais-177-joao-pessoa. Acesso em: 30 out. 2026.'),
    ]
    for lead, title, tail in references:
        add_reference(doc, lead, title, tail)

    # Ensure headings/captions and table labels don't inherit an unwanted blue theme.
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.color and run.font.color.rgb:
                run.font.color.rgb = None

    doc.core_properties.title = TITLE
    doc.core_properties.subject = 'Desenvolvimento e avaliação técnica — Prioridade 2'
    doc.core_properties.author = 'Arthur Nobrega Martins; Hygor Gabriel Alves Montenegro; Thiago Rodrigues Medeiros'
    doc.save(str(OUTPUT))

    words_pt = re.findall(r"\b[\wÀ-ÿ]+(?:[-’'][\wÀ-ÿ]+)*\b", ABSTRACT_PT)
    words_en = re.findall(r"\b[\wÀ-ÿ]+(?:[-’'][\wÀ-ÿ]+)*\b", ABSTRACT_EN)
    if not (150 <= len(words_pt) <= 250 and 150 <= len(words_en) <= 250):
        raise SystemExit(f'Abstract fora da faixa: PT={len(words_pt)}, EN={len(words_en)}')
    print(f'output={OUTPUT}')
    print(f'abstract_words_pt={len(words_pt)} abstract_words_en={len(words_en)}')
    print(f'figures={len([p for p in doc.inline_shapes])} tables={len(doc.tables)} paragraphs={len(doc.paragraphs)}')


if __name__ == '__main__':
    main()
